"""Starting a question set from a file instead of typing it.

Covers the reading (services/question_import.py), Word files
(services/document_service.py), and the panel on both pages putting what it
read into the question table.
"""

from __future__ import annotations

import io
import json

import pandas as pd
import pytest

from services import llm_service
from services import question_import as importer
from services.document_service import extract_text
from tests.test_form_page_fixes import _app_as, _button, _labels
from utils.validation import validate_upload

def _texts(result):
    return [q.question_text for q in result.questions]


# ---------------------------------------------------------------------------
# Reading the questions out of the text
# ---------------------------------------------------------------------------
def test_numbered_questions_are_split_with_their_marks():
    result = importer.import_questions(
        "Algebra worksheet 4\nAnswer all questions.\n\n"
        "1. Solve 2x + 3 = 11. [2]\n"
        "2) Solve 5 - x = 2.\n   Show your working. (3 marks)\n"
        "3. Expand 2(x + 3)."
    )

    assert result.method == importer.METHOD_NUMBERS
    assert _texts(result) == [
        "Solve 2x + 3 = 11.",
        "Solve 5 - x = 2.\nShow your working.",
        "Expand 2(x + 3).",
    ]
    assert [q.max_marks for q in result.questions] == [2, 3, 1]
    assert [q.marks_found for q in result.questions] == [True, True, False]
    assert "Marks were read for 2" in result.message


def test_parts_stay_in_their_question_and_their_marks_add_up():
    result = importer.import_questions(
        "1. A number is multiplied by 4 and then 6 is subtracted. The result is 26.\n"
        "(a) Form an equation. [1]\n"
        "(b) Solve it. [2]\n"
        "2(a) Write 0.25 as a fraction. [1]\n"
        "2(b) Write 3/4 as a percentage. [1]\n"
        "3. Simplify 18 : 21. [Total: 2]"
    )

    assert len(result.questions) == 3
    first, second, third = result.questions
    assert "(a) Form an equation. [1]" in first.question_text
    assert "(b) Solve it. [2]" in first.question_text
    assert first.max_marks == 3
    assert second.question_text.startswith("(a) Write 0.25")
    assert "2(b) Write 3/4" in second.question_text
    assert second.max_marks == 2
    assert third.max_marks == 2 and "Total" not in third.question_text


def test_a_numbered_list_inside_a_question_is_not_split():
    result = importer.import_questions(
        "1. Solve the equation 3x + 1 = 10 using these steps:\n"
        "1. subtract 1 from both sides\n"
        "2. divide both sides by 3\n"
        "2. Expand 2(x + 3). [2]"
    )

    assert len(result.questions) == 2
    assert "2. divide both sides by 3" in result.questions[0].question_text
    assert result.questions[1].question_text == "Expand 2(x + 3)."


def test_data_lines_are_not_mistaken_for_question_numbers():
    result = importer.import_questions(
        "1. Find the mean of these numbers:\n3 5 7 9\n2. Find their range."
    )

    assert _texts(result) == ["Find the mean of these numbers:\n3 5 7 9", "Find their range."]


def test_a_list_of_steps_does_not_swallow_the_next_questions():
    result = importer.import_questions(
        "1. Expand 2(x+1). [1]\n"
        "2. Solve 3x + 1 = 10 using these steps:\n1. subtract 1\n2. divide by 3\n"
        "3. Factorise x^2 + 5x + 6. [2]\n"
        "4. Simplify 18 : 21. [1]"
    )

    assert len(result.questions) == 4
    assert result.questions[1].question_text.endswith("2. divide by 3")
    assert [q.max_marks for q in result.questions] == [1, 1, 2, 1]


def test_one_numbered_question_with_parts_stays_one_question():
    result = importer.import_questions(
        "Rectangle worksheet\n1. A rectangle is 6 cm by 4 cm.\n"
        "(a) Find its area. [2]\n(b) Find its perimeter. [2]"
    )

    assert result.method == importer.METHOD_NUMBERS
    assert len(result.questions) == 1
    assert result.questions[0].max_marks == 4


def test_long_runs_of_spaces_are_read_in_linear_time():
    import time

    started = time.perf_counter()
    importer.import_questions("Q" + " " * 20_000 + "!\nAnswer" + " " * 20_000 + "!\n1. A\n2. B")
    assert time.perf_counter() - started < 1.0


def test_printed_papers_with_bare_numbers_and_answer_lines():
    result = importer.import_questions(
        "1 Work out 3.5 x 2. ............................................ [1]\n"
        "2 (a) Write 0.25 as a fraction. ..................... [1]\n"
        "(b) Write 3/4 as a percentage. ....................... [1]\n"
        "x = ...........................\n"
        "© UCLES 2024 0580/12/M/J/24 [Turn over\n"
    )

    assert _texts(result) == [
        "Work out 3.5 x 2.",
        "(a) Write 0.25 as a fraction. [1]\n(b) Write 3/4 as a percentage. [1]",
    ]
    assert [q.max_marks for q in result.questions] == [1, 2]


def test_answers_inline_and_in_an_answer_section():
    result = importer.import_questions(
        "Question 1\nWork out 3 + 4 x 2.\nAnswer: 11\n"
        "Question 2\nSimplify 18 : 21 [1]\n"
        "Question 3\nWrite 0.5 as a percentage.\n\n"
        "Answers\n1. 11\n2. 6 : 7\n3. 50%"
    )

    assert [q.model_answer for q in result.questions] == ["11", "6 : 7", "50%"]
    assert result.answers_found_count == 3


def test_unnumbered_text_falls_back_to_lines_or_paragraphs():
    lines = importer.import_questions("What is 7 x 8?\nWrite 0.5 as a percentage.")
    assert lines.method == importer.METHOD_LINES
    assert _texts(lines) == ["What is 7 x 8?", "Write 0.5 as a percentage."]
    assert "Delete any that are instructions" in lines.message

    paragraphs = importer.import_questions("Find x if\n2x = 8.\n\nFind y if\ny + 1 = 3.")
    assert paragraphs.method == importer.METHOD_PARAGRAPHS
    assert len(paragraphs.questions) == 2


def test_nothing_is_invented_and_empty_text_finds_nothing():
    result = importer.import_questions("1. Solve 2x = 8.\n2. Solve 3x = 9.")
    assert all(q.model_answer == "" for q in result.questions)
    assert all(q.max_marks == importer.DEFAULT_MARKS for q in result.questions)
    assert "starts at 1 mark" in result.message

    assert not importer.import_questions("   ").found


def test_editor_rows_and_suggested_title():
    result = importer.import_questions("1. Solve 2x = 8. [2]\n2. Solve 3x = 9.")
    assert result.editor_rows() == [
        {"question_text": "Solve 2x = 8.", "model_answer": "", "marking_criteria": "", "max_marks": 2.0},
        {"question_text": "Solve 3x = 9.", "model_answer": "", "marking_criteria": "", "max_marks": 1.0},
    ]
    assert importer.suggested_title("algebra_worksheet-4.pdf") == "Algebra worksheet 4"
    assert importer.suggested_title("") == ""


# ---------------------------------------------------------------------------
# The AI reading keeps only what is in the file
# ---------------------------------------------------------------------------
class _FakeModel:
    name = "fake"

    def __init__(self, reply: dict) -> None:
        self.reply = reply
        self.calls = 0

    def complete(self, system: str, user: str) -> str:
        self.calls += 1
        return json.dumps(self.reply)


@pytest.fixture(autouse=True)
def _empty_ai_cache():
    llm_service._CACHE.clear()
    yield
    llm_service._CACHE.clear()


def test_ai_reading_keeps_only_the_files_own_text():
    """The AI only points at questions: what is saved is the file's text."""
    source = (
        "Worksheet\n"
        "1. Solve 2x + 3 = 11 for x. [2]\n"
        "2. Work out the area of a\n"
        "rectangle 3 cm by 4 cm.\n"
        "Answers\n"
        "1. x = 4\n"
    )
    model = _FakeModel({"questions": [
        {"question": "Solve 2x + 3 = 11 for x.", "answer": "x = 4", "marks": 2},
        {"question": "Work out the area of a rectangle 3 cm by 4 cm.", "answer": "12 cm²", "marks": 5},
        # Every word is in the file, but it is not the file's question.
        {"question": "Solve 2x - 3 = 11 for x.", "answer": "", "marks": None},
        {"question": "Prove that the square root of 2 is irrational.", "answer": "", "marks": 4},
        {"question": "x = 4", "answer": "", "marks": None},
    ]})

    result = importer.ai_import_questions(source, model=model)

    assert result.method == importer.METHOD_AI
    assert _texts(result) == [
        "Solve 2x + 3 = 11 for x.",
        "Work out the area of a\nrectangle 3 cm by 4 cm.",
    ]
    first, second = result.questions
    # The mark printed after the question on its line, and the answer key's
    # entry for question 1 - both the file's.
    assert (first.model_answer, first.max_marks) == ("x = 4", 2)
    # The AI's "12 cm²" and 5 marks are in the file nowhere: never used.
    assert (second.model_answer, second.max_marks, second.marks_found) == ("", 1, False)
    assert "3 suggestions were left out" in result.message


def test_ai_reading_is_cached_and_none_without_a_model(monkeypatch):
    model = _FakeModel({"questions": [{"question": "Solve 2x = 8.", "answer": "", "marks": None}]})
    importer.ai_import_questions("1. Solve 2x = 8.", model=model)
    importer.ai_import_questions("1. Solve 2x = 8.", model=model)
    assert model.calls == 1

    monkeypatch.setattr(llm_service, "get_model", lambda settings=None: None)
    assert importer.ai_import_questions("1. Solve 2x = 8.") is None


# ---------------------------------------------------------------------------
# Word files
# ---------------------------------------------------------------------------
def _docx(build) -> bytes:
    import docx

    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_a_word_numbered_list_keeps_its_numbers():
    def build(document):
        document.add_paragraph("Algebra worksheet")
        for text in ("Solve 2x + 3 = 11. [2]", "Solve 5 - x = 2. (3 marks)", "Expand 2(x + 3)."):
            document.add_paragraph(text, style="List Number")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "x"
        table.rows[0].cells[1].text = "y"

    extraction = extract_text("worksheet.docx", _docx(build))

    assert extraction.success, extraction.message
    assert "1. Solve 2x + 3 = 11. [2]" in extraction.text
    assert "3. Expand 2(x + 3)." in extraction.text
    assert "x | y" in extraction.text
    result = importer.import_questions(extraction.text)
    assert [q.max_marks for q in result.questions] == [2, 3, 1]


def test_a_file_named_docx_must_be_a_word_document():
    ok, message = validate_upload("paper.docx", b"%PDF-1.7 not a word file")
    assert not ok and "not a Word document" in message
    ok, message = validate_upload("paper.docx", b"PK\x03\x04 not really a zip")
    assert not ok and "not a Word document" in message
    empty = _docx(lambda document: None)
    assert not extract_text("empty.docx", empty).success


# ---------------------------------------------------------------------------
# The panel puts the questions in the table
# ---------------------------------------------------------------------------
PAPER = "1. Solve 2x + 3 = 11. [2]\n2. Solve 5 - x = 2.\nAnswers\n1. x = 4\n2. x = 3"


def test_my_questions_imports_a_paper_into_the_table():
    at = _app_as("student")
    at.switch_page("pages/my_questions.py")
    at.session_state["mq_import_text"] = PAPER
    at.session_state["mq_import_filename"] = "Linear_equations_sheet.pdf"
    at.run()
    assert not at.exception, at.exception

    _button(at, "Put these 2 questions in the table").click().run()

    assert not at.exception, at.exception
    rows = at.session_state["my_questions_rows"]
    assert isinstance(rows, pd.DataFrame)
    assert list(rows["question_text"]) == ["Solve 2x + 3 = 11.", "Solve 5 - x = 2."]
    assert list(rows["model_answer"]) == ["x = 4", "x = 3"]
    assert list(rows["max_marks"]) == [2.0, 1.0]
    assert at.text_input(key="mq_title").value == "Linear equations sheet"
    assert any("Put 2 questions in the table" in s.value for s in at.success)


def test_create_assessment_imports_a_paper_into_the_table():
    at = _app_as("teacher")
    at.switch_page("pages/create_assessment.py")
    at.session_state["ca_import_text"] = PAPER
    at.run()
    assert not at.exception, at.exception
    assert "Put these 2 questions in the table" in _labels(at)

    _button(at, "Put these 2 questions in the table").click().run()

    assert not at.exception, at.exception
    rows = at.session_state["create_assessment_rows"]
    assert list(rows.columns) == ["question_text", "model_answer", "marking_criteria", "max_marks"]
    assert len(rows) == 2


def test_a_word_list_can_start_from_its_own_number():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def build(document):
        for text in ("Solve 2x = 8.", "Solve 3x = 9."):
            document.add_paragraph(text, style="List Number")
        # Point both paragraphs at a new list that starts at 5.
        numbering = document.part.numbering_part.element
        # The decimal list behind the "List Number" style.
        style_num_id = str(document.styles["List Number"].element.pPr.numPr.numId.val)
        style_num = next(
            n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == style_num_id
        )
        abstract_id = style_num.find(qn("w:abstractNumId")).get(qn("w:val"))
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), "77")
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract_id)
        num.append(ref)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "5")
        override.append(start)
        num.append(override)
        numbering.append(num)
        for paragraph in document.paragraphs:
            ppr = paragraph._p.get_or_add_pPr()
            num_pr = ppr.get_or_add_numPr()
            num_pr.get_or_add_numId().val = 77
            num_pr.get_or_add_ilvl().val = 0

    extraction = extract_text("later.docx", _docx(build))

    assert extraction.success, extraction.message
    assert "5. Solve 2x = 8." in extraction.text
    assert "6. Solve 3x = 9." in extraction.text


def test_an_impossible_word_list_level_is_ignored():
    def build(document):
        paragraph = document.add_paragraph("Solve 2x = 8 for x please.", style="List Number")
        paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_ilvl().val = 1_000_000

    extraction = extract_text("odd.docx", _docx(build))

    assert extraction.success, extraction.message
    assert extraction.text == "Solve 2x = 8 for x please."


def test_a_zip_with_an_unsupported_version_is_not_a_word_document():
    data = bytearray(_docx(lambda document: document.add_paragraph("Solve 2x = 8 for x please.")))
    central = data.find(b"PK\x01\x02")
    data[central + 6:central + 8] = (255).to_bytes(2, "little")  # version needed to extract

    ok, message = validate_upload("odd.docx", bytes(data))

    assert not ok and "not a Word document" in message
    assert not extract_text("odd.docx", bytes(data)).success
